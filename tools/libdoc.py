#!/usr/bin/env python3
"""libdoc.py spec.json out.docx   — पूर्वोदाहरण-ग्रंथालयासाठी Word (सर्व मजकूर फॉन्ट DVOT-SurekhMR).
spec = {"title": "...", "subtitle": "...", "sections": [
   {"h": "शीर्षक", "level": 1,                # level 1 = मोठे, 2 = उप-शीर्षक, 3 = लहान
    "paras": ["..."],                          # साधे परिच्छेद
    "bullets": ["..."],                        # • यादी
    "quotes": [{"text": "जसाच्या तसा उतारा", "src": "स्रोत (पत्र क्र./दि./पृ./प्रकरण)"}],
    "table": {"cols": ["अ", "ब"], "rows": [["..", ".."]], "widths_cm": [3, 12]},   # optional
    "pagebreak": false } ] }
"""
import sys, json
import docx
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

FONT = "DVOT-SurekhMR"

def setfont(run, size=None, bold=None, italic=None):
    rpr = run._element.get_or_add_rPr(); rf = rpr.get_or_add_rFonts()
    for a in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"): rf.set(qn(a), FONT)
    if size: run.font.size = Pt(size)
    if bold is not None: run.bold = bold
    if italic is not None: run.italic = italic

def newdoc():
    d = docx.Document()
    s = d.sections[0]
    s.page_width, s.page_height = Cm(21.0), Cm(29.7)
    s.left_margin = s.right_margin = Cm(2.0); s.top_margin = s.bottom_margin = Cm(1.8)
    for name in ("Normal", "Table Grid", "List Bullet"):
        try:
            st = d.styles[name]; st.font.name = FONT; st.font.size = Pt(12)
            rf = st.element.get_or_add_rPr().get_or_add_rFonts()
            for a in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"): rf.set(qn(a), FONT)
        except Exception: pass
    return d

def para(d, text, size=12, bold=False, italic=False, align=None, left=None, space_after=4):
    p = d.add_paragraph(); r = p.add_run(text); setfont(r, size, bold, italic)
    pf = p.paragraph_format; pf.space_after = Pt(space_after); pf.line_spacing = 1.15
    if align == "center": p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    elif align == "justify": p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    if left is not None: pf.left_indent = Cm(left)
    return p

def shade(cell, color="D9D9D9"):
    tcPr = cell._element.get_or_add_tcPr(); sh = OxmlElement("w:shd")
    sh.set(qn("w:val"), "clear"); sh.set(qn("w:color"), "auto"); sh.set(qn("w:fill"), color); tcPr.append(sh)

def table(d, cols, rows, widths=None):
    t = d.add_table(rows=1, cols=len(cols)); t.style = "Table Grid"
    for i, c in enumerate(cols):
        cell = t.rows[0].cells[i]; cell.text = ""; r = cell.paragraphs[0].add_run(str(c)); setfont(r, 11, True); shade(cell)
    for row in rows:
        cells = t.add_row().cells
        for i, v in enumerate(row):
            cells[i].text = ""; r = cells[i].paragraphs[0].add_run(str(v)); setfont(r, 11)
    if widths:
        for row in t.rows:
            for i, w in enumerate(widths): row.cells[i].width = Cm(w)
    # repeat header row
    trPr = t.rows[0]._tr.get_or_add_trPr(); h = OxmlElement("w:tblHeader"); h.set(qn("w:val"), "true"); trPr.append(h)
    d.add_paragraph()
    return t

def build(spec, out):
    d = newdoc()
    para(d, spec["title"], 18, True, align="center", space_after=6)
    if spec.get("subtitle"): para(d, spec["subtitle"], 12, False, align="center", space_after=10)
    for sec in spec["sections"]:
        if sec.get("pagebreak"):
            d.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
        lvl = sec.get("level", 1)
        if sec.get("h"): para(d, sec["h"], {1: 15, 2: 13, 3: 12}.get(lvl, 12), True, space_after=4)
        for t in sec.get("paras", []): para(d, t, align="justify")
        for t in sec.get("bullets", []): para(d, "• " + t, left=0.6, align="justify")
        for q in sec.get("quotes", []):
            para(d, "“" + q["text"] + "”", 11.5, False, left=1.0, align="justify", space_after=1)
            if q.get("src"): para(d, "(स्रोत: " + q["src"] + ")", 10.5, False, True, left=1.0, space_after=6)
        if sec.get("table"):
            tb = sec["table"]; table(d, tb["cols"], tb["rows"], tb.get("widths_cm"))
    d.save(out)

if __name__ == "__main__":
    build(json.load(open(sys.argv[1], encoding="utf-8")), sys.argv[2]); print("ok", sys.argv[2])
