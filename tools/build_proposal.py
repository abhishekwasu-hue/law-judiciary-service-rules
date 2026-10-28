#!/usr/bin/env python3
"""प्रस्ताव-पॅकेज: नोंद (docx) + पत्र (docx) + जोडपत्र PDF -> एकच printable PDF.
Usage: build_proposal.py out.pdf --note note.docx [--letter letter.docx] \
        --enc "pdf/path.pdf|जोडपत्र-अ: शीर्षक" [--enc ...]
Each enclosure gets a Devanagari separator page (जोडपत्र क्रमांक + शीर्षक + मूळ PDF पान-संख्या);
the original PDF pages are appended UNCHANGED.
"""
import argparse, subprocess, tempfile, os, sys
from pypdf import PdfReader, PdfWriter
from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH

def to_pdf(docx, outdir):
    subprocess.run(["soffice","--headless","--convert-to","pdf","--outdir",outdir,docx],check=True,capture_output=True)
    return os.path.join(outdir, os.path.splitext(os.path.basename(docx))[0]+".pdf")

def separator(title, pages, tmp, i):
    d = Document()
    for _ in range(8): d.add_paragraph()
    for text, size in ((title, 24), (f"(मूळ दस्तऐवज — {pages} पाने, जसाच्या तसा)", 14)):
        p = d.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(text); r.font.size = Pt(size); r.font.name = "Mangal"; r.bold = size > 20
    f = os.path.join(tmp, f"sep{i}.docx"); d.save(f); return to_pdf(f, tmp)

def main():
    a = argparse.ArgumentParser()
    a.add_argument("out"); a.add_argument("--note"); a.add_argument("--letter")
    a.add_argument("--enc", action="append", default=[])
    o = a.parse_args(); w = PdfWriter(); tmp = tempfile.mkdtemp()
    for f in (o.note, o.letter):
        if f:
            p = f if f.endswith(".pdf") else to_pdf(f, tmp)
            for pg in PdfReader(p).pages: w.add_page(pg)
    for i, e in enumerate(o.enc, 1):
        path, title = e.split("|", 1)
        r = PdfReader(path)
        for pg in PdfReader(separator(title, len(r.pages), tmp, i)).pages: w.add_page(pg)
        for pg in r.pages: w.add_page(pg)
    w.write(o.out); print("written", o.out, len(w.pages), "pages")
main()
